"""Exports a summary as .txt or .docx (FR4, UR8, NFR7). Files are built in memory only, never saved (NFR3)."""

import io

from docx import Document

import config


def to_txt(summary) -> bytes:
    return summary.as_text().encode("utf-8")


def to_docx(summary) -> bytes:
    doc = Document()
    doc.add_heading(summary.title, level=1)
    for line in summary.body.splitlines():
        text = line.strip().replace("**", "")
        if not text:
            continue
        if text.startswith("#"):
            doc.add_heading(text.lstrip("#").strip(), level=2)
        elif text.startswith(("- ", "* ")):
            doc.add_paragraph(text[2:], style="List Bullet")
        else:
            doc.add_paragraph(text)
    if summary.source_url:
        doc.add_paragraph(f"Original lecture: {summary.source_url}")
    doc.add_paragraph().add_run(config.AI_DISCLAIMER).italic = True  # NFR12

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()