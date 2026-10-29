"""Unit tests for summarizer/export.py (the .txt and .docx downloads)."""

import io

from docx import Document

import config 

from summarizer.export import to_docx, to_txt
from summarizer.summarize import Summary

SUMMARY = Summary(
    title="Lecture 3",
    body="## Main topics\n- Unit testing\n- **pytest** basics\nA short closing sentence.",
)


def test_to_txt_contains_title_and_body():
    text = to_txt(SUMMARY).decode("utf-8")
    assert text.startswith("Lecture 3")
    assert "- Unit testing" in text


def test_to_docx_is_a_real_word_file_with_the_right_content():
    data = to_docx(SUMMARY)
    doc = Document(io.BytesIO(data)) 
    paragraphs = [p.text for p in doc.paragraphs]

    assert paragraphs[0] == "Lecture 3"
    assert "Main topics" in paragraphs
    assert "pytest basics" in paragraphs  
    assert doc.paragraphs[2].style.name == "List Bullet"

def test_exports_include_the_ai_disclaimer():
    assert config.AI_DISCLAIMER in to_txt(SUMMARY).decode("utf-8")

    doc = Document(io.BytesIO(to_docx(SUMMARY)))
    last = doc.paragraphs[-1]
    assert last.text == config.AI_DISCLAIMER
    assert last.runs[0].italic


def test_no_original_lecture_line_in_any_export():
    # SCRUM-41: the link to the original lecture was removed.
    assert "Original lecture" not in to_txt(SUMMARY).decode("utf-8")
    doc = Document(io.BytesIO(to_docx(SUMMARY)))
    assert not any("Original lecture" in p.text for p in doc.paragraphs)


def test_custom_title_is_used_in_both_exports():
    summary = Summary(title="My own title", body="- point")
    assert to_txt(summary).decode("utf-8").startswith("My own title")
    assert Document(io.BytesIO(to_docx(summary))).paragraphs[0].text == "My own title"
